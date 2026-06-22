import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    # Initialize presentation
    prs = Presentation()
    
    # Custom Theme Colors
    primary_color = RGBColor(13, 71, 161)    # Deep Blue
    secondary_color = RGBColor(21, 101, 192) # Medium Blue
    text_color = RGBColor(33, 33, 33)        # Dark Gray

    # Slide 1: Title Slide
    slide_layout = prs.slide_layouts[0] # Title slide
    slide1 = prs.slides.add_slide(slide_layout)
    title1 = slide1.shapes.title
    subtitle1 = slide1.placeholders[1]

    title1.text = "Zanzibar Metropolitan College (ZMC)"
    title1.text_frame.paragraphs[0].font.color.rgb = primary_color
    title1.text_frame.paragraphs[0].font.bold = True
    
    subtitle1.text = "Learning Management System\nComprehensive Platform for Students, Instructors, and Administrators"
    subtitle1.text_frame.paragraphs[0].font.color.rgb = secondary_color

    # Slide 2: System Overview
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide2 = prs.slides.add_slide(slide_layout)
    title2 = slide2.shapes.title
    content2 = slide2.placeholders[1]

    title2.text = "System Overview & Architecture"
    title2.text_frame.paragraphs[0].font.color.rgb = primary_color
    
    tf2 = content2.text_frame
    tf2.text = "Core Modules & Technology Stack:"
    p = tf2.add_paragraph()
    p.text = "• Modern Next.js Frontend with an interactive, responsive UI"
    p.level = 1
    p = tf2.add_paragraph()
    p.text = "• Robust Laravel Backend for secure data management & API"
    p.level = 1
    p = tf2.add_paragraph()
    p.text = "• Role-Based Access Control (Admin, Instructor, Student, Accountant)"
    p.level = 1
    p = tf2.add_paragraph()
    p.text = "• Complete lifecycle management from registration to graduation"
    p.level = 1

    # Slide 3: Student Experience
    slide3 = prs.slides.add_slide(slide_layout)
    title3 = slide3.shapes.title
    content3 = slide3.placeholders[1]

    title3.text = "Empowering the Student Journey"
    title3.text_frame.paragraphs[0].font.color.rgb = primary_color

    tf3 = content3.text_frame
    tf3.text = "A Premium, User-Centric Portal:"
    p = tf3.add_paragraph()
    p.text = "• Seamless multi-step registration with dynamic fee calculations"
    p.level = 1
    p = tf3.add_paragraph()
    p.text = "• Real-time access to academic records, timetables, and exam results"
    p.level = 1
    p = tf3.add_paragraph()
    p.text = "• Integrated panels for finance, accommodation, and secure settings"
    p.level = 1

    # Slide 4: Instructor & Admin Capabilities
    slide4 = prs.slides.add_slide(slide_layout)
    title4 = slide4.shapes.title
    content4 = slide4.placeholders[1]

    title4.text = "Management & Academic Control"
    title4.text_frame.paragraphs[0].font.color.rgb = primary_color

    tf4 = content4.text_frame
    tf4.text = "Tools for Staff and Administrators:"
    p = tf4.add_paragraph()
    p.text = "• Instructors can easily grade students, manage results, and monitor course progress via bulk uploads."
    p.level = 1
    p = tf4.add_paragraph()
    p.text = "• Admins have full oversight of users, programs, departments, and financial statistics."
    p.level = 1
    p = tf4.add_paragraph()
    p.text = "• Dynamic GPA regulations and automated academic standing (Good Standing vs. Discontinued)."
    p.level = 1

    # Slide 5: Conclusion & Future Outlook
    slide5 = prs.slides.add_slide(slide_layout)
    title5 = slide5.shapes.title
    content5 = slide5.placeholders[1]

    title5.text = "Transforming Education Through Tech"
    title5.text_frame.paragraphs[0].font.color.rgb = primary_color

    tf5 = content5.text_frame
    tf5.text = "Looking Forward:"
    p = tf5.add_paragraph()
    p.text = "• Automated report generation and deep data analytics"
    p.level = 1
    p = tf5.add_paragraph()
    p.text = "• Scalable infrastructure designed to support growing student enrollments"
    p.level = 1
    p = tf5.add_paragraph()
    p.text = "• Ensuring data integrity, absolute security, and a premium user experience"
    p.level = 1

    # Save the presentation
    filename = "ZMC_LMS_Presentation.pptx"
    prs.save(filename)
    print(f"Presentation saved successfully as {filename}")

if __name__ == "__main__":
    create_presentation()
